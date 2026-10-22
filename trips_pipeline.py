"""Build the Wave ETP trips report step by step, saving a file after every step for debugging.

Output goes to "Steps/":
  01 Combined.xlsx         every row from every input file, with Site Name and File Name
  02 Deduplicated.xlsx     duplicate rows removed (SN and File Name ignored)
  03 Sorted.xlsx           sorted by Vehicle Name, then Entry Time (Exit Time when there's no entry)
  04 Trip Rows Marked.xlsx Trip column + debug columns showing the next row and why each row did/didn't match
  05 Trip Rows Only.xlsx   only rows marked "Trip Row"
  06 Trip Details.xlsx     Trip Number, Vehicle, Start/End Time, Start/End Site, Duration filled on start rows
  07 Trips.xlsx            one formatted row per trip
"""

import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
INPUT_DIR = BASE / "Input Files"
STEPS_DIR = BASE / "Steps"
COLUMNS = ["SN", "Entry Time", "Exit Time", "Vehicle Name", "Duration (dd:hh:mm)"]
IGNORE = {"7_3_2026Wave ETP.xls"}  # hand-built cross-site sheet, not a raw export
WAVE = "Wave ETP"


def save(df: pd.DataFrame, name: str, formats: dict[str, str] | None = None) -> None:
    """Write a step file with readable widths, bold header, filter and frozen header row."""
    path = STEPS_DIR / name
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="dd-mmm-yyyy hh:mm:ss") as writer:
        df.to_excel(writer, index=False)
        ws = writer.sheets["Sheet1"]
        for i, col in enumerate(df.columns, start=1):
            letter = get_column_letter(i)
            values = [str(col)] + [str(v) for v in df[col].head(200)]
            ws.column_dimensions[letter].width = min(max(len(v) for v in values) + 2, 45)
            if formats and col in formats:
                for cell in ws[letter][1:]:
                    cell.number_format = formats[col]
                    cell.alignment = Alignment(horizontal="center")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    print(f"{name}: {len(df)} rows")


# ---------- Step 1: combine ----------

def site_from_filename(name: str) -> str:
    stem = Path(name).stem
    stem = re.sub(r"^\d{1,2}_\d{1,2}_\d{4}", "", stem)  # leading date, e.g. 10_15_2026
    stem = re.sub(r"\s*\(\d+\)$", "", stem)  # trailing copy suffix, e.g. (1)
    return stem.strip()


def read_file(path: Path) -> pd.DataFrame:
    raw = pd.read_excel(path, header=None, dtype=object)
    # Header row isn't always the first row (some exports have a "Site" line above it)
    header_idx = next(
        (i for i, row in raw.iterrows() if {"SN", "Vehicle Name"} <= set(row.astype(str).str.strip())),
        None,
    )
    if header_idx is None:
        raise ValueError(f"{path.name}: no header row with 'SN' and 'Vehicle Name' found")
    df = raw.iloc[header_idx + 1 :].copy()
    df.columns = raw.iloc[header_idx].astype(str).str.strip()
    df = df[COLUMNS]
    # Data rows have a numeric SN; this drops blank lines and any footer text
    return df[pd.to_numeric(df["SN"], errors="coerce").notna()]


def to_datetime(series: pd.Series) -> pd.Series:
    series = series.replace("", None)
    # Cells formatted as dates arrive as datetimes; unformatted ones as Excel serial numbers
    is_num = series.map(lambda v: isinstance(v, (int, float)) and not isinstance(v, bool))
    from_serial = pd.to_datetime(pd.to_numeric(series.where(is_num)), unit="D", origin="1899-12-30")
    as_dates = pd.to_datetime(series.where(~is_num), errors="coerce")
    return as_dates.fillna(from_serial).dt.round("s")


def combine() -> pd.DataFrame:
    frames = []
    for path in sorted(INPUT_DIR.glob("*.xls*")):
        if path.name in IGNORE:
            print(f"  skipped {path.name}")
            continue
        df = read_file(path)
        df["Site Name"] = site_from_filename(path.name)
        df["File Name"] = path.name
        frames.append(df)
    if not frames:
        raise FileNotFoundError(f"No input Excel files found in {INPUT_DIR}")
    df = pd.concat(frames, ignore_index=True)
    df["SN"] = pd.to_numeric(df["SN"]).astype(int)
    df["Entry Time"] = to_datetime(df["Entry Time"])
    df["Exit Time"] = to_datetime(df["Exit Time"])
    df["Vehicle Name"] = df["Vehicle Name"].astype(str).str.strip()
    return df


# ---------- Steps 2-3: dedup and sort ----------

def dedupe(df: pd.DataFrame) -> pd.DataFrame:
    """Drop repeated events. SN is only the row number within an export and File Name differs
    between exports, so the same event in two exports differs in both - ignore them."""
    return df.drop_duplicates(subset=df.columns.drop(["SN", "File Name"]), ignore_index=True)


def sort_rows(df: pd.DataFrame) -> pd.DataFrame:
    """By vehicle, then entry time. Rows with no entry time (visit started before the export
    window) are placed by their exit time instead of being pushed to the end."""
    sort_time = df["Entry Time"].fillna(df["Exit Time"])
    return (
        df.assign(_sort_time=sort_time)
        .sort_values(["Vehicle Name", "_sort_time"], kind="stable", ignore_index=True)
        .drop(columns="_sort_time")
    )


# ---------- Step 4: mark trip rows ----------

def mark_trips(df: pd.DataFrame) -> pd.DataFrame:
    """A trip = a Wave ETP row with an exit time, followed by a non-Wave ETP row with an entry time."""
    df = df.copy()
    nxt = df.shift(-1)
    is_wave = df["Site Name"] == WAVE
    trip_start = is_wave & df["Exit Time"].notna()
    trip_end = ~is_wave & df["Entry Time"].notna()
    pair = trip_start & trip_end.shift(-1, fill_value=False)
    df["Trip"] = ""
    df.loc[pair | pair.shift(1, fill_value=False), "Trip"] = "Trip Row"

    # Debug columns: what the next row looks like and why a Wave ETP row did or didn't start a trip
    df["Next Vehicle"] = nxt["Vehicle Name"]
    df["Next Site"] = nxt["Site Name"]
    df["Next Entry Time"] = nxt["Entry Time"]
    df["Same Vehicle As Next"] = df["Vehicle Name"] == nxt["Vehicle Name"]
    reason = pd.Series("", index=df.index)
    reason[is_wave & df["Exit Time"].isna()] = "Wave ETP row has no exit time"
    reason[trip_start & (nxt["Site Name"] == WAVE)] = "next row is also Wave ETP"
    reason[trip_start & (nxt["Site Name"] != WAVE) & nxt["Entry Time"].isna()] = "next row has no entry time"
    reason[trip_start & (df["Vehicle Name"] != nxt["Vehicle Name"])] = "last row for this vehicle"
    reason[trip_start & nxt["Site Name"].isna()] = "last row in file"
    reason[pair] = "trip start"
    reason[pair & (df["Vehicle Name"] != nxt["Vehicle Name"])] = "trip start - but next row is a DIFFERENT vehicle"
    reason[pair.shift(1, fill_value=False)] = "trip end"
    df["Debug Reason"] = reason
    return df


# ---------- Step 6: trip details ----------

def trip_details(df: pd.DataFrame) -> pd.DataFrame:
    df = df.reset_index(drop=True)
    nxt = df.shift(-1)
    is_start = (df["Site Name"] == WAVE) & df["Exit Time"].notna() & (df["Vehicle Name"] == nxt["Vehicle Name"])
    trip_no = is_start.cumsum()
    df["Trip Number"] = trip_no.where(is_start | is_start.shift(1, fill_value=False)).astype("Int64")
    df.loc[is_start, "Vehicle"] = df["Vehicle Name"]
    df.loc[is_start, "Start Time"] = df["Exit Time"]
    df.loc[is_start, "End Time"] = nxt["Entry Time"]
    df.loc[is_start, "Start Site"] = df["Site Name"]
    df.loc[is_start, "End Site"] = nxt["Site Name"]
    df["Duration"] = df["End Time"] - df["Start Time"]
    return df


def main():
    STEPS_DIR.mkdir(exist_ok=True)

    df = combine()
    save(df, "01 Combined.xlsx")

    before = len(df)
    df = dedupe(df)
    save(df, "02 Deduplicated.xlsx")
    print(f"  dedup: {before} -> {len(df)} rows")

    df = sort_rows(df)
    save(df, "03 Sorted.xlsx")

    df = mark_trips(df)
    save(df, "04 Trip Rows Marked.xlsx")

    df = df[df["Trip"] == "Trip Row"].drop(
        columns=["Next Vehicle", "Next Site", "Next Entry Time", "Same Vehicle As Next", "Debug Reason"]
    )
    save(df, "05 Trip Rows Only.xlsx")

    df = trip_details(df)
    save(df, "06 Trip Details.xlsx", {"Duration": "[h]:mm:ss"})

    trips = df[df["Vehicle"].notna()][["Trip Number", "Vehicle", "Start Time", "End Time", "Start Site", "End Site", "Duration"]]
    save(trips, "07 Trips.xlsx", {"Trip Number": "0", "Duration": "[h]:mm:ss"})


if __name__ == "__main__":
    main()
